#!/usr/bin/env python3
"""Build Academic CV (4–6 pp) and Academic Bio (1 pp) in English and Chinese."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public" / "cv"

NAVY = HexColor("#0b1f33")
WETLAND = HexColor("#1c4a3c")
INK = HexColor("#12202b")
MUTED = HexColor("#3a4650")

pdfmetrics.registerFont(TTFont("LibSerif", "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"))
pdfmetrics.registerFont(TTFont("LibSerif-Bold", "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"))
pdfmetrics.registerFont(TTFont("WQY", "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"))

CONTACT_EN = (
    "Email: kpyi@rcees.ac.cn · Office: 010-62915655 · "
    "18 Shuangqing Road, Haidian District, Beijing 100085 · "
    "https://dse.rcees.cas.cn/ · "
    "https://yikunpeng82-sys.github.io/kunpeng-yi-website/"
)
CONTACT_ZH = (
    "邮箱：kpyi@rcees.ac.cn · 办公电话：010-62915655 · "
    "北京海淀区双清路18号 100085 · "
    "https://dse.rcees.cas.cn/ · "
    "https://yikunpeng82-sys.github.io/kunpeng-yi-website/"
)

TIMELINE_EN = [
    "Associate Professor, Research Center for Eco-Environmental Sciences, Chinese Academy of Sciences; State Key Laboratory of Regional and Urban Ecology (SKLRUE), formerly State Key Laboratory of Urban and Regional Ecology — April 2017 – present",
    "Postdoctoral Fellow (Ecological Remote Sensing), Aerospace Information Research Institute, Chinese Academy of Sciences (AIRCAS), formerly RADI — August 2014 – April 2017",
    "Ph.D., Remote Sensing, Hokkaido University — October 2010 – June 2014. Thesis: 遥感卫星观测及模型在生物质火烧中的应用研究",
    "M.Sc., Physical Geography, Northeast Normal University — September 2008 – July 2010",
    "B.Sc., Geography, Harbin College — September 2003 – July 2007",
]

TIMELINE_ZH = [
    "副研究员，中国科学院生态环境研究中心、区域与城市生态安全全国重点实验室 — 2017年4月至今",
    "博士后（生态遥感），中国科学院空天信息创新研究院（AIRCAS，原遥感与数字地球研究所 RADI） — 2014年8月–2017年4月",
    "遥感科学博士，北海道大学 — 2010年10月–2014年6月。学位论文：遥感卫星观测及模型在生物质火烧中的应用研究",
    "自然地理学硕士，东北师范大学 — 2008年9月–2010年7月",
    "地理学学士，哈尔滨学院 — 2003年9月–2007年7月",
]

RESEARCH_EN = (
    "Research addresses global climate change and eco-environmental remote sensing. "
    "Typical ecosystems—wetlands, forests and grasslands—in cities, nature reserves and national parks "
    "are examined in relation to drought, flood, wildfire and other extreme meteorological hazards. "
    "Threatened migratory birds are treated as indicator species for migration mechanisms, flyway-corridor "
    "conservation, and responses to human activity and extreme climate events. Applications include Sino BON, "
    "restoration of typical degraded wetlands in Northeast China, bird-habitat planning for the Xiong'an New Area, "
    "Nanchang territorial spatial planning (2019–2035), and a bird-strike risk early-warning system. "
    "Public academic identity: Movement Ecology × Earth Observation × Biodiversity Conservation."
)

OVERVIEW_ZH = (
    "日本北海道大学博士，中国科学院空天信息创新研究院（原遥感与数字地球研究所）博士后，"
    "现为中国科学院生态环境研究中心副研究员，任职于区域与城市生态安全全国重点实验室。"
    "主要从事全球气候变化与生态环境遥感监测研究。以城市、自然保护区、国家公园内的湿地、森林、草地等典型生态系统为目标，"
    "研究干旱、洪涝、野火等极端气象灾害对生态系统的影响机制。以珍稀濒危的迁徙鸟类为生态系统的指示物种，"
    "探究鸟类迁徙机制、候鸟迁飞通道保护策略、迁徙鸟类对人类活动和极端气候事件响应和反馈机制。"
    "相关成果服务于 Sino BON、东北典型退化湿地恢复与重建技术及示范、雄安新区鸟类栖息地规划、"
    "南昌市国土空间规划（2019-2035）及鸟撞风险预警系统等。"
    "主持国家自然科学基金、国家重点研发计划子课题、地方委托等科研项目10余项，主持项目总经费800余万元，"
    "发表学术论文40余篇，发明专利3项，参编专著2部。"
)

SERVICE_EN = (
    "Council member, China Grassland Society. Member of the Geographical Society of China, "
    "the Ecological Society of China, the Chinese Society of Remote Sensing, and the China Ornithological Society. "
    "Guest Editor of special issues for Remote Sensing, Land and Fire."
)
SERVICE_ZH = (
    "中国草学会理事；中国地理学会、中国生态学会、中国遥感学会、中国鸟类学会会员；"
    "Remote Sensing、Land、Fire 专刊 Guest Editor。"
)

NSFC_EN = [
    "鸟类多尺度运动行为模式及其迁徙路线核心栖息地监测与保护研究. NSFC General Programme, 2023.01–2026.12, Principal investigator",
    "基于多源遥感和卫星追踪大数据的鸟类运动行为信息反演研究. NSFC General Programme, 2019.01–2022.12, Principal investigator",
]
NSFC_ZH = [
    "鸟类多尺度运动行为模式及其迁徙路线核心栖息地监测与保护研究。国家自然科学基金面上项目，2023.01–2026.12，主持",
    "基于多源遥感和卫星追踪大数据的鸟类运动行为信息反演研究。国家自然科学基金面上项目，2019.01–2022.12，主持",
]

STUDENTS_EN = [
    "Master’s supervisor (硕导).",
    "Directions: 运动生态学，迁徙动物监测与保护; 湿地监测与生态修复; 生态系统遥感监测与评估.",
    "Majors: 071300 Ecology; 0705Z1 Natural Resources Science; 070503 Cartography and Geographic Information Systems.",
    "Prospective students: write to kpyi@rcees.ac.cn.",
]
STUDENTS_ZH = [
    "硕导。",
    "招生方向：运动生态学，迁徙动物监测与保护；湿地监测与生态修复；生态系统遥感监测与评估。",
    "招生专业：071300生态学；0705Z1自然资源学；070503地图学与地理信息系统。",
    "意向学生请写信至 kpyi@rcees.ac.cn。",
]

OUTPUT_EN = (
    "More than 40 academic papers; more than ten hosted projects (National Natural Science Foundation of China, "
    "a National Key R&D Programme sub-project, and local commissions) with total hosted funding exceeding RMB 8 million; "
    "three invention patents; two co-authored or edited books."
)
OUTPUT_ZH = "发表学术论文40余篇；主持科研项目10余项，主持经费800余万元；发明专利3项；参编专著2部。"

SELECTED_EN = [
    "Zhang, X., Wan, W., Yang, H. & Yi, K. (2025). Rapid development of wind energy infrastructure threatens bird migration. The Innovation Geoscience 3(3): 100142. https://doi.org/10.59717/j.xinn-geo.2025.100142",
    "Tian, Z., Huo, D., Yi, K., Que, J., Lu, Z. & Hou, J. (2024). Evaluation of Suitable Habitats for Birds Based on MaxEnt and Google Earth Engine—A Case Study of Baer’s Pochard (Aythya baeri) in Baiyangdian, China. Remote Sensing 16(1): 64. https://doi.org/10.3390/rs16010064",
    "Yi, K., Meng, F., Gu, D. & Miao, Q. (2023). Optimizing Water Level Management Strategies to Strengthen Reservoir Support for Bird’s Migration Network. Remote Sensing 15(23): 5508. https://doi.org/10.3390/rs15235508",
    "Yi, K., Zhang, J., Batbayar, N., Higuchi, H., Natsagdorj, T. & Bysykatova, I. P. (2022). Using Tracking Data to Identify Gaps in Knowledge and Conservation of the Critically Endangered Siberian Crane (Leucogeranus leucogeranus). Remote Sensing 14(20): 5101. https://doi.org/10.3390/rs14205101",
    "Batbayar, N., Yi, K., Zhang, J., Natsagdorj, T., Damba, I., Cao, L. & Fox, A. D. (2021). Combining Tracking and Remote Sensing to Identify Critical Year-Round Site, Habitat Use and Migratory Connectivity of a Threatened Waterbird Species. Remote Sensing 13(20): 4049. https://doi.org/10.3390/rs13204049",
]

PAPERS = [
    "Yu, M., Yin, L. & Yi, K. (2026). Bird diversity responses to mega-events influence: A case study of the Beijing 2022 Winter Olympics. Ecological Informatics 93: 103594. https://doi.org/10.1016/j.ecoinf.2025.103594",
    "Zhang, X., Wan, W., Yang, H. & Yi, K. (2025). Rapid development of wind energy infrastructure threatens bird migration. The Innovation Geoscience 3(3): 100142. https://doi.org/10.59717/j.xinn-geo.2025.100142",
    "Yin, S., Yi, K., Zhang, X., Nie, T., Meng, L., Sun, Z., Chu, Q., Ai, Z., Zhao, X., Wu, L., Guo, M. & Liu, X. (2024). Temporal and Spatial Dynamics of Summer Crop Residue Burning Practices in North China. Remote Sensing 16(24): 4763. https://doi.org/10.3390/rs16244763",
    "Tian, Z., Huo, D., Yi, K., Que, J., Lu, Z. & Hou, J. (2024). Evaluation of Suitable Habitats for Birds Based on MaxEnt and Google Earth Engine—A Case Study of Baer’s Pochard (Aythya baeri) in Baiyangdian, China. Remote Sensing 16(1): 64. https://doi.org/10.3390/rs16010064",
    "Yi, K., Meng, F., Gu, D. & Miao, Q. (2023). Optimizing Water Level Management Strategies to Strengthen Reservoir Support for Bird’s Migration Network. Remote Sensing 15(23): 5508. https://doi.org/10.3390/rs15235508",
    "Yi, K., Zhao, X., Zheng, Z., Zhao, D. & Zeng, Y. (2023). Trends of greening and browning in terrestrial vegetation in China from 2000 to 2020. Ecological Indicators 154: 110587. https://doi.org/10.1016/j.ecolind.2023.110587",
    "Yi, K., Zhang, J., Batbayar, N., Higuchi, H., Natsagdorj, T. & Bysykatova, I. P. (2022). Using Tracking Data to Identify Gaps in Knowledge and Conservation of the Critically Endangered Siberian Crane (Leucogeranus leucogeranus). Remote Sensing 14(20): 5101. https://doi.org/10.3390/rs14205101",
    "Bao, Y., Shinoda, M., Yi, K., et al. (2022). Satellite-Based Analysis of Spatiotemporal Wildfire Pattern in the Mongolian Plateau. Remote Sensing 15(1): 190. https://doi.org/10.3390/rs15010190",
    "Ren, X., Yi, K. & Cao, L. (2022). 优化风电场和电网布局以避免鸟撞. 环境保护科学 48(3): 13–19.",
    "Batbayar, N., Yi, K., Zhang, J., Natsagdorj, T., Damba, I., Cao, L. & Fox, A. D. (2021). Combining Tracking and Remote Sensing to Identify Critical Year-Round Site, Habitat Use and Migratory Connectivity of a Threatened Waterbird Species. Remote Sensing 13(20): 4049. https://doi.org/10.3390/rs13204049",
    "Damba, I., Zhang, J., Yi, K., et al. (2021). Seasonal and regional differences in migration patterns and conservation status of Swan Geese (Anser cygnoides) in the East Asian Flyway. Avian Research 12(1): 73. https://doi.org/10.1186/s40657-021-00308-y",
    "Li, H., Fang, L., Wang, X., Yi, K., Cao, L. & Fox, A. D. (2020). Does snowmelt constrain spring migration progression in sympatric wintering Arctic-nesting geese? Ibis 162(2): 548–555. https://doi.org/10.1111/ibi.12767",
    "Zhu, Q., Damba, I., Zhao, Q., Yi, K., et al. (2020). Lack of conspicuous sex-biased dispersal patterns at different spatial scales in an Asian endemic goose species breeding in unpredictable steppe wetlands. Ecology and Evolution 10: 7006–7020. https://doi.org/10.1002/ece3.6382",
    "Damba, I., Fang, L., Yi, K., et al. (2020). Flyway structure, breeding, migration and wintering distributions of the globally threatened Swan Goose Anser cygnoides in East Asia. Wildfowl: 97–123.",
    "Yan, M., Yi, K., Zhang, J., et al. (2020). Flyway connectivity and population status of the Greylag Goose Anser anser in East Asia. Wildfowl 6: 157–180.",
    "Mao, D., Wang, Z., Wu, J., Wu, B., Zeng, Y., Song, K., Yi, K. & Luo, L. (2018). China's wetlands loss to urban expansion. Land Degradation & Development 29(8): 2644–2657. https://doi.org/10.1002/ldr.2939",
    "Wang, X., Cao, L., Bysykatova, I., et al. (2018). The Far East taiga forest: unrecognized inhospitable terrain for migrating Arctic-nesting waterbirds? PeerJ 6: e4353. https://doi.org/10.7717/peerj.4353",
    "Yi, K., Bao, Y. & Zhang, J. (2017). Spatial distribution and temporal variability of open fire in China. International Journal of Wildland Fire 26(2): 122–135. https://doi.org/10.1071/WF15213",
    "Yi, K., Zeng, Y. & Wu, B. (2016). Mapping and evaluation the process, pattern and potential of urban growth in China. Applied Geography 71: 44–55. https://doi.org/10.1016/j.apgeog.2016.04.011",
    "Yi, K. & Bao, Y. (2016). Estimates of Wildfire Emissions in Boreal Forests of China. Forests 7(8): 158. https://doi.org/10.3390/f7080158",
    "Yi, K., Tani, H., Li, Q., Zhang, J., Guo, M., Bao, Y., Wang, X. & Li, J. (2014). Mapping and Evaluating the Urbanization Process in Northeast China Using DMSP/OLS Nighttime Light Data. Sensors 14(2): 3207–3226. https://doi.org/10.3390/s140203207",
    "Yi, K., Tani, H., Zhang, J., Guo, M., Wang, X. & Zhong, G. (2013). Long-Term Satellite Detection of Post-Fire Vegetation Trends in Boreal Forests of China. Remote Sensing 5(12): 6938–6957. https://doi.org/10.3390/rs5126938",
    "Guo, M., Wang, X., Li, J., Yi, K., Zhong, G. & Tani, H. (2012). Assessment of Global Carbon Dioxide Concentration Using MODIS and GOSAT Data. Sensors 12(12): 16368–16389. https://doi.org/10.3390/s121216368",
]


def styles(cjk: bool, compact: bool = False):
    base = getSampleStyleSheet()
    font = "WQY"
    body_size = 10 if compact else (11 if cjk else 11)
    item_size = 9 if compact else (10.4 if cjk else 10.6)
    leading = 14 if compact else (18 if cjk else 16.2)
    return {
        "title": ParagraphStyle(
            "T", parent=base["Title"], fontName="WQY", fontSize=16 if compact else 18,
            textColor=NAVY, alignment=TA_LEFT, spaceAfter=2 * mm, leading=20,
        ),
        "sub": ParagraphStyle(
            "S", parent=base["Normal"], fontName="WQY", fontSize=10,
            textColor=WETLAND, spaceAfter=3 * mm, leading=14,
        ),
        "h": ParagraphStyle(
            "H", parent=base["Heading2"], fontName="WQY", fontSize=11 if compact else 12.5,
            textColor=NAVY, spaceBefore=3.5 * mm if compact else 6 * mm, spaceAfter=2 * mm, leading=16,
        ),
        "body": ParagraphStyle(
            "B", parent=base["Normal"], fontName=font, fontSize=body_size,
            textColor=INK, alignment=TA_JUSTIFY, leading=leading, spaceAfter=1.6 * mm,
        ),
        "meta": ParagraphStyle(
            "M", parent=base["Normal"], fontName=font, fontSize=body_size,
            textColor=INK, leading=leading, spaceAfter=0.8 * mm,
        ),
        "item": ParagraphStyle(
            "I", parent=base["Normal"], fontName="WQY", fontSize=item_size,
            textColor=INK, leading=leading - 1,
        ),
    }


def bullets(items: list[str], style: ParagraphStyle) -> ListFlowable:
    return ListFlowable(
        [ListItem(Paragraph(item, style), leftIndent=8) for item in items],
        bulletType="bullet",
        start="•",
        leftIndent=12,
        bulletFontName="WQY",
        bulletFontSize=style.fontSize,
    )


def footer_cv(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(MUTED)
    canvas.setFont("WQY", 8)
    canvas.drawString(16 * mm, 10 * mm, "Kunpeng Yi / 伊坤朋 — Academic CV / 学术简历")
    canvas.drawRightString(A4[0] - 16 * mm, 10 * mm, f"{doc.page}")
    canvas.restoreState()


def footer_bio(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(MUTED)
    canvas.setFont("WQY", 8)
    canvas.drawString(16 * mm, 10 * mm, "Kunpeng Yi / 伊坤朋 — Academic Bio / 学术简介")
    canvas.drawRightString(A4[0] - 16 * mm, 10 * mm, f"{doc.page}")
    canvas.restoreState()


def build_pdf(path: Path, story, title: str, author: str, footer, compact: bool = False) -> None:
    doc = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm,
        topMargin=16 * mm, bottomMargin=18 * mm,
        title=title, author=author,
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def build_en_cv(path: Path) -> None:
    s = styles(False)
    story = [
        Paragraph("Academic CV", s["title"]),
        Paragraph("Kunpeng Yi / 伊坤朋", s["sub"]),
        Paragraph("Associate Professor, Research Center for Eco-Environmental Sciences, Chinese Academy of Sciences", s["meta"]),
        Paragraph("State Key Laboratory of Regional and Urban Ecology (SKLRUE)", s["meta"]),
        Paragraph(CONTACT_EN, s["meta"]),
        Paragraph("Education and appointments", s["h"]),
        bullets(TIMELINE_EN, s["item"]),
        Paragraph("Research", s["h"]),
        Paragraph("Research direction: 生态环境遥感研究.", s["body"]),
        Paragraph(RESEARCH_EN, s["body"]),
        Paragraph("Hosted projects", s["h"]),
        bullets(NSFC_EN, s["item"]),
        Paragraph(
            "Also hosted a National Key R&D Programme sub-project and locally commissioned studies. "
            "Output overview: " + OUTPUT_EN,
            s["body"],
        ),
        Paragraph("Selected publications", s["h"]),
        bullets(SELECTED_EN, s["item"]),
        Paragraph("Publications", s["h"]),
        bullets(PAPERS, s["item"]),
        Paragraph("Patents and books", s["h"]),
        Paragraph("Three invention patents. Two co-authored or edited books.", s["body"]),
        Paragraph("Professional service", s["h"]),
        Paragraph(SERVICE_EN, s["body"]),
        Paragraph("Students", s["h"]),
        bullets(STUDENTS_EN, s["item"]),
        Paragraph("Contact", s["h"]),
        Paragraph(CONTACT_EN, s["body"]),
        Spacer(1, 8 * mm),
        Paragraph(
            "Laboratory English name: State Key Laboratory of Regional and Urban Ecology (SKLRUE), "
            "formerly State Key Laboratory of Urban and Regional Ecology. "
            "Laboratory website: https://dse.rcees.cas.cn/",
            s["body"],
        ),
    ]
    build_pdf(path, story, "Academic CV — Kunpeng Yi", "Kunpeng Yi", footer_cv)


def build_zh_cv(path: Path) -> None:
    s = styles(True)
    story = [
        Paragraph("学术简历", s["title"]),
        Paragraph("伊坤朋（Kunpeng Yi）", s["sub"]),
        Paragraph("中国科学院生态环境研究中心副研究员", s["meta"]),
        Paragraph("区域与城市生态安全全国重点实验室", s["meta"]),
        Paragraph(CONTACT_ZH, s["meta"]),
        Paragraph("学历与任职", s["h"]),
        bullets(TIMELINE_ZH, s["item"]),
        Paragraph("研究", s["h"]),
        Paragraph("研究方向：生态环境遥感研究。", s["body"]),
        Paragraph(OVERVIEW_ZH, s["body"]),
        Paragraph("主持项目", s["h"]),
        bullets(NSFC_ZH, s["item"]),
        Paragraph("另主持国家重点研发计划子课题及地方委托项目。", s["body"]),
        Paragraph("代表性论文", s["h"]),
        bullets(SELECTED_EN, s["item"]),
        Paragraph("论文", s["h"]),
        bullets(PAPERS, s["item"]),
        Paragraph("专利与著作", s["h"]),
        Paragraph("发明专利3项；参编专著2部。", s["body"]),
        Paragraph("学术任职", s["h"]),
        Paragraph(SERVICE_ZH, s["body"]),
        Paragraph("招生", s["h"]),
        bullets(STUDENTS_ZH, s["item"]),
        Paragraph("联系", s["h"]),
        Paragraph(CONTACT_ZH, s["body"]),
    ]
    build_pdf(path, story, "学术简历 — 伊坤朋", "伊坤朋", footer_cv)


def build_en_bio(path: Path) -> None:
    s = styles(False, compact=True)
    story = [
        Paragraph("Academic Bio", s["title"]),
        Paragraph("Kunpeng Yi / 伊坤朋", s["sub"]),
        Paragraph("Associate Professor, Research Center for Eco-Environmental Sciences, Chinese Academy of Sciences", s["meta"]),
        Paragraph("State Key Laboratory of Regional and Urban Ecology (SKLRUE)", s["meta"]),
        Paragraph(CONTACT_EN, s["meta"]),
        Paragraph("Profile", s["h"]),
        Paragraph(RESEARCH_EN, s["body"]),
        Paragraph(OUTPUT_EN, s["body"]),
        Paragraph("Appointments", s["h"]),
        bullets(TIMELINE_EN[:3], s["item"]),
        Paragraph("Selected publications", s["h"]),
        bullets(SELECTED_EN[:4], s["item"]),
        Paragraph("Students and service", s["h"]),
        Paragraph(SERVICE_EN, s["body"]),
        bullets(STUDENTS_EN, s["item"]),
    ]
    build_pdf(path, story, "Academic Bio — Kunpeng Yi", "Kunpeng Yi", footer_bio, compact=True)


def build_zh_bio(path: Path) -> None:
    s = styles(True, compact=True)
    story = [
        Paragraph("学术简介", s["title"]),
        Paragraph("伊坤朋（Kunpeng Yi）", s["sub"]),
        Paragraph("中国科学院生态环境研究中心副研究员", s["meta"]),
        Paragraph("区域与城市生态安全全国重点实验室", s["meta"]),
        Paragraph(CONTACT_ZH, s["meta"]),
        Paragraph("简介", s["h"]),
        Paragraph(OVERVIEW_ZH, s["body"]),
        Paragraph("学历与任职", s["h"]),
        bullets(TIMELINE_ZH[:3], s["item"]),
        Paragraph("代表性论文", s["h"]),
        bullets(SELECTED_EN[:4], s["item"]),
        Paragraph("招生与任职", s["h"]),
        Paragraph(SERVICE_ZH, s["body"]),
        bullets(STUDENTS_ZH, s["item"]),
    ]
    build_pdf(path, story, "学术简介 — 伊坤朋", "伊坤朋", footer_bio, compact=True)


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x0B, 0x1F, 0x33)


def build_en_docx(path: Path) -> None:
    doc = Document()
    title = doc.add_paragraph()
    run = title.add_run("Academic CV — Kunpeng Yi / 伊坤朋")
    run.bold = True
    run.font.size = Pt(18)
    doc.add_paragraph("Associate Professor, Research Center for Eco-Environmental Sciences, Chinese Academy of Sciences")
    doc.add_paragraph("State Key Laboratory of Regional and Urban Ecology (SKLRUE)")
    doc.add_paragraph(CONTACT_EN)
    add_heading(doc, "Education and appointments")
    for item in TIMELINE_EN:
        doc.add_paragraph(item, style="List Bullet")
    add_heading(doc, "Research")
    doc.add_paragraph(RESEARCH_EN)
    add_heading(doc, "Hosted projects")
    for item in NSFC_EN:
        doc.add_paragraph(item, style="List Bullet")
    add_heading(doc, "Selected publications")
    for paper in SELECTED_EN:
        doc.add_paragraph(paper, style="List Bullet")
    add_heading(doc, "Publications")
    for paper in PAPERS:
        doc.add_paragraph(paper, style="List Bullet")
    add_heading(doc, "Patents and books")
    doc.add_paragraph("Three invention patents. Two co-authored or edited books.")
    add_heading(doc, "Professional service")
    doc.add_paragraph(SERVICE_EN)
    add_heading(doc, "Students")
    for item in STUDENTS_EN:
        doc.add_paragraph(item, style="List Bullet")
    doc.save(str(path))


def build_zh_docx(path: Path) -> None:
    doc = Document()
    title = doc.add_paragraph()
    run = title.add_run("学术简历 — 伊坤朋（Kunpeng Yi）")
    run.bold = True
    run.font.size = Pt(18)
    doc.add_paragraph("中国科学院生态环境研究中心副研究员")
    doc.add_paragraph("区域与城市生态安全全国重点实验室")
    doc.add_paragraph(CONTACT_ZH)
    add_heading(doc, "学历与任职")
    for item in TIMELINE_ZH:
        doc.add_paragraph(item, style="List Bullet")
    add_heading(doc, "研究")
    doc.add_paragraph(OVERVIEW_ZH)
    add_heading(doc, "主持项目")
    for item in NSFC_ZH:
        doc.add_paragraph(item, style="List Bullet")
    add_heading(doc, "代表性论文")
    for paper in SELECTED_EN:
        doc.add_paragraph(paper, style="List Bullet")
    add_heading(doc, "论文")
    for paper in PAPERS:
        doc.add_paragraph(paper, style="List Bullet")
    add_heading(doc, "专利与著作")
    doc.add_paragraph("发明专利3项；参编专著2部。")
    add_heading(doc, "学术任职")
    doc.add_paragraph(SERVICE_ZH)
    add_heading(doc, "招生")
    for item in STUDENTS_ZH:
        doc.add_paragraph(item, style="List Bullet")
    doc.save(str(path))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    build_en_cv(OUT / "kunpeng-yi-cv-en.pdf")
    build_zh_cv(OUT / "kunpeng-yi-cv-zh.pdf")
    build_en_bio(OUT / "kunpeng-yi-bio-en.pdf")
    build_zh_bio(OUT / "kunpeng-yi-bio-zh.pdf")
    build_en_docx(OUT / "kunpeng-yi-cv-en.docx")
    build_zh_docx(OUT / "kunpeng-yi-cv-zh.docx")
    for name in sorted(p.name for p in OUT.iterdir()):
        p = OUT / name
        print(f"{p.name:32} {p.stat().st_size:8} bytes")


if __name__ == "__main__":
    main()
